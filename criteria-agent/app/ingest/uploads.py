"""Upload handling."""

import asyncio
import base64
import csv
import io
import zipfile
from collections import defaultdict
from functools import cache
from pathlib import Path

import httpx
import numpy
import openpyxl
import pypdfium2
from lxml import etree
from openpyxl.worksheet.properties import PageSetupProperties
from opentelemetry.trace import Span
from python_calamine import CalamineWorkbook
from rapid_orientation import RapidOrientation

from app.ai.jev import banking, likely
from app.ai.llm import ask_model, bounded
from app.core import database
from app.core.config import client, settings
from app.core.documents import PDFIUM, Doc, edition, forget, normalize, opened, png
from app.core.errors import BLANK, DAMAGED, LOCKED, NO_OCR, NOT_BANKING, UNPARSED, Unusable
from app.core.events import emit
from app.core.prompts import UPRIGHT
from app.core.schemas import Upload, Upright
from app.core.tracing import failure, observed, trace_id, tracer
from app.ingest.ocr import ocr
from app.ingest.structure import ensure, structure

SHEETS = {".xlsx", ".xlsm", ".csv"}
LOCKS: defaultdict[str, asyncio.Lock] = defaultdict(asyncio.Lock)
LEGACY = {".xls", ".xlsb", ".ods"}
WORD = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
REJECTED = " | ".join(
    [
        "//w:tr[w:trPr/w:del]",
        "//w:del",
        "//w:moveFrom",
        *(
            f"//w:{change}"
            for change in (
                "rPrChange",
                "pPrChange",
                "tblPrChange",
                "trPrChange",
                "tcPrChange",
                "sectPrChange",
                "tblGridChange",
                "numberingChange",
            )
        ),
    ]
)


def converter() -> httpx.AsyncClient:
    """The converter client."""
    return client(base_url=settings().converter_url)


def workbook(raw: bytes, name: str) -> openpyxl.Workbook:
    """A spreadsheet upload as a workbook; CSV and older formats bring their cell values."""
    suffix = Path(name).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return openpyxl.load_workbook(io.BytesIO(raw))
    if suffix == ".csv":
        sheets = {"Sheet1": list(csv.reader(io.StringIO(raw.decode("utf-8-sig", errors="replace"))))}
    else:
        read = CalamineWorkbook.from_filelike(io.BytesIO(raw))
        sheets = {title: read.get_sheet_by_name(title).to_python() for title in read.sheet_names}
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append([None if value == "" else value for value in row])
    return book


def whole(pdf: bytes, raw: bytes) -> bool:
    """Whether a converted spreadsheet prints every text cell in full."""
    printed = normalize("".join(Doc("", "", pdf).layer))
    read = CalamineWorkbook.from_filelike(io.BytesIO(raw))
    rows = (row for title in read.sheet_names for row in read.get_sheet_by_name(title).to_python())
    return all(normalize(value) in printed for row in rows for value in row if isinstance(value, str))


def fitted(book: openpyxl.Workbook) -> bytes:
    """A workbook whose columns fit their text, each sheet printed one page wide."""
    for sheet in book.worksheets:
        for column in sheet.iter_cols():
            longest = max((len(str(cell.value)) for cell in column if cell.value is not None), default=0)
            sheet.column_dimensions[column[0].column_letter].width = longest + 1
        sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@cache
def orientation() -> RapidOrientation:
    """The page orientation classifier, on one thread so it never takes every core."""
    return RapidOrientation(cfg_path=Path(__file__).with_name("orientation.yaml"))


def guessed(pdf: bytes, layer: tuple[str, ...]) -> dict[int, tuple[int, bytes]]:
    """Scanned pages the orientation classifier would turn, holding the PDF lock only to read and render."""
    found = {}
    scans = [n for n, text in enumerate(layer) if not text.strip()]
    with PDFIUM:
        document = pypdfium2.PdfDocument(pdf)
    try:
        for number in scans:
            with PDFIUM:
                image = document[number].render().to_pil().convert("RGB")
            label, _ = orientation()(numpy.asarray(image))
            if int(label):
                found[number] = (int(label), png(image.rotate(int(label), expand=True)))
    finally:
        with PDFIUM:
            document.close()
    return found


def turned(pdf: bytes, turns: dict[int, int]) -> bytes:
    """The PDF with these pages turned counter-clockwise by these angles."""
    if not turns:
        return pdf
    with PDFIUM, pypdfium2.PdfDocument(pdf) as document:
        for number, turn in turns.items():
            page = document[number]
            page.set_rotation((page.get_rotation() - turn) % 360)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()


async def upright(pdf: bytes, layer: tuple[str, ...]) -> bytes:
    """The PDF with each scanned page turned upright where a vision check confirms the classifier's turn."""
    guesses = await asyncio.to_thread(guessed, pdf, layer)
    views = await bounded(
        [
            lambda image=image: ask_model(UPRIGHT, "Look at the image.", Upright, images=(image,))
            for _, image in guesses.values()
        ]
    )
    turns = {number: turn for (number, (turn, _)), view in zip(guesses.items(), views, strict=True) if view.upright}
    return await asyncio.to_thread(turned, pdf, turns)


async def portable(raw: bytes, name: str) -> bytes:
    """The upload as a PDF."""
    if raw.startswith(b"%PDF-"):
        return raw
    emit("ingest", stage="convert")
    suffix, flat = Path(name).suffix.lower(), f"{Path(name).stem}.xlsx"
    if suffix in SHEETS:
        return await converted(await asyncio.to_thread(lambda: fitted(workbook(raw, name))), flat)
    if suffix in (".docx", ".docm") and zipfile.is_zipfile(io.BytesIO(raw)):
        raw = await asyncio.to_thread(accepted, raw)
    pdf = await converted(raw, name)
    if suffix in LEGACY and not await asyncio.to_thread(whole, pdf, raw):
        return await converted(await asyncio.to_thread(lambda: fitted(workbook(raw, name))), flat)
    return pdf


def accepted(raw: bytes) -> bytes:
    """A Word file with its tracked changes accepted, as Word shows it without markup."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item)
            if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                tree = etree.fromstring(data)
                for gone in tree.xpath(REJECTED, namespaces=WORD):
                    gone.getparent().remove(gone)
                for kept in tree.xpath("//w:ins | //w:moveTo", namespaces=WORD):
                    parent, at = kept.getparent(), kept.getparent().index(kept)
                    parent[at : at + 1] = list(kept)
                data = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(item, data)
    return buffer.getvalue()


async def converted(raw: bytes, name: str) -> bytes:
    """A file converted to PDF."""
    with tracer.start_as_current_span("gotenberg.convert") as span:
        async with converter() as client:
            response = await client.post("/forms/libreoffice/convert", files={"files": (name, raw)})
        observed(
            span,
            {"file": name, "bytes": len(raw)},
            {"status": response.status_code, "pdf bytes": len(response.content)},
        )
    if response.is_client_error:
        raise Unusable(f"{name} could not be converted to PDF.")
    return response.raise_for_status().content


async def parsing(pdf: bytes, pages: list[int], layer: tuple[str, ...]) -> dict[int, dict]:
    """OCR pages if available."""
    if not settings().mistral_api_key or not pages:
        return {}
    emit("ingest", stage="parse", pages=len(pages))
    try:
        read = await ocr(pdf, pages, layer)
    except Exception as error:
        emit("ingest", stage="parse", status="failed", reason=f"{UNPARSED} ({type(error).__name__})")
        return {}
    emit("ingest", stage="parse", status="done", pages=len(pages))
    return {p["number"] - 1: p for p in read}


async def opening(pdf: bytes, layer: tuple[str, ...]) -> tuple[str, dict[int, dict]]:
    """The document's opening words."""
    words, read = [], {}
    for number, text in enumerate(layer):
        if not text.strip():
            read |= await parsing(pdf, [number], layer)
            text = read.get(number, {}).get("text", "")
        words += text.split()
        if len(words) >= settings().excerpt_words:
            break
    return " ".join(words[: settings().excerpt_words]), read


async def receive(upload: Upload, session: str) -> dict:
    """Handle an upload."""
    attributes = {"session.id": session, "langfuse.trace.tags": ["chat", "upload"]}
    with tracer.start_as_current_span("upload", attributes=attributes) as span:
        observed(span, {"file": upload.name, "base64 characters": len(upload.content)})
        try:
            pdf = await portable(base64.b64decode(upload.content, validate=True), upload.name)
            id_ = edition(pdf)
            async with LOCKS[id_]:
                response = await admitted(upload.name, id_, pdf, span)
        except Exception as error:
            return failure(span, error)
        span.set_attribute("langfuse.observation.output", response.get("document_id", NOT_BANKING))
        return {**response, "name": upload.name, "trace_id": trace_id()}


def unreadable(where: str) -> Unusable:
    """Why no text was found: a blank page, or a scan with no OCR set up to read it."""
    return Unusable((BLANK if settings().mistral_api_key else NO_OCR).format(where=where))


async def admitted(name: str, id_: str, pdf: bytes, span: Span) -> dict:
    """Accept or reject an upload, keyed by the file as uploaded; a new scan is turned upright first."""
    try:
        layer = await asyncio.to_thread(lambda: Doc(id_, name, pdf).layer)
    except pypdfium2.PdfiumError as error:
        locked = error.err_code == pypdfium2.raw.FPDF_ERR_PASSWORD
        raise Unusable((LOCKED if locked else DAMAGED).format(name=name)) from None
    stored = await database.known(id_)
    if stored or (stored is not None and not settings().mistral_api_key):
        await database.save(id_, name, pdf, None)
        await ensure(await opened(id_))
        emit("ingest", stage="stored", pages=len(layer), parsed=stored)
        return {"status": "uploaded", "document_id": id_, "pages": len(layer), "parsed": stored}
    pdf = await upright(pdf, layer)
    excerpt, read = await opening(pdf, layer)
    if not excerpt:
        raise unreadable("the first page")
    emit("ingest", stage="classify")
    verdict = await banking(excerpt)
    emit(
        "ingest",
        stage="classify",
        status="done",
        banking=likely(verdict.noul),
        confidence=verdict.noul,
        excerpt=excerpt,
    )
    span.set_attribute("upload.banking", likely(verdict.noul))
    if not likely(verdict.noul):
        return {"status": "rejected", "reason": NOT_BANKING}
    count = len(layer)
    read |= await parsing(pdf, [n for n in range(count) if n not in read], layer)
    for number in range(count):
        emit("ingest", stage="read", page=number + 1, pages=count)
    if not any(layer[n].strip() or read.get(n, {}).get("text", "").strip() for n in range(count)):
        raise unreadable("any page")
    parsed = [read[n] for n in range(count)] if len(read) == count else None
    await database.save(id_, name, pdf, parsed)
    forget(id_)
    await structure(await opened(id_))
    emit("ingest", stage="stored", pages=count, parsed=parsed is not None)
    return {"status": "uploaded", "document_id": id_, "pages": count, "parsed": parsed is not None}
